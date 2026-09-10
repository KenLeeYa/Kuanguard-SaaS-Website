import csv
import base64
from datetime import datetime, timedelta
from io import BytesIO, StringIO
from pathlib import PurePosixPath
import re
import zipfile
from typing import Literal

from fastapi import APIRouter, Depends, Query, Request
from pydantic import Field, field_validator

from . import models as m, wallet
from .config import settings
from .db import add, all_rows, aware, change, one
from .learning_routes import assign, course_row
from .project_routes import Input
from .security import Context, context, digest, fail, idempotent, owned, paged

router = APIRouter()


class RecipientInput(Input):
    name: str = Field(min_length=1, max_length=120)
    csv: str = Field(min_length=1, max_length=2_000_000)


class RecipientWorkbookInput(Input):
    name: str = Field(min_length=1, max_length=120)
    filename: str = Field(pattern=r"^[^/\\:]+\.xlsx$")
    content_base64: str = Field(min_length=1, max_length=2_800_000)


class CampaignInput(Input):
    name: str = Field(min_length=1, max_length=120)
    group_id: str
    scheduled_at: datetime
    remediation_course_id: str | None = None

    @field_validator("scheduled_at")
    @classmethod
    def timezone_required(cls, value):
        if value.tzinfo is None:
            raise ValueError("Timezone required")
        return value


class ConfirmationInput(Input):
    confirmed: Literal[True]


class DispatchInput(Input):
    outcome: Literal["accepted", "rejected", "unknown"] = "accepted"


class SuspiciousInput(Input):
    campaign_id: str | None = None
    description: str = Field(min_length=1, max_length=5000)


class EventReviewInput(Input):
    classification: Literal["confirmed_human", "automated", "ambiguous"]
    reason: str = Field(min_length=10, max_length=2000)


class ReconciliationInput(Input):
    outcome: Literal["accepted", "rejected"]
    provider_reference: str = Field(min_length=1, max_length=120)
    occurred_at: datetime
    reason: str = Field(min_length=10, max_length=2000)

    @field_validator("occurred_at")
    @classmethod
    def verified_time(cls, value):
        if value.tzinfo is None or value > m.now():
            raise ValueError("Confirmed event time requires a timezone and cannot be in the future")
        return value


def owned_campaign(ctx, campaign_id):
    row = owned(ctx, m.campaigns, campaign_id)
    if row["owner_id"] != ctx.user_id:
        fail(404, "CAMPAIGN_SCOPE_DENIED", "找不到您獲授權管理的活動。")
    return row


def campaign_view(ctx, row):
    messages = all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.campaign_id == row["id"])
    events = all_rows(ctx.conn, m.raw_events, ctx.tenant_id, m.raw_events.c.campaign_id == row["id"])
    metrics = {"planned": len(messages), "provider_accepted": sum(message["status"] in {"accepted", "bounced"} for message in messages),
               "delivered": len({event["message_id"] for event in events if event["event_type"] == "delivered"}),
               "bounced": len({event["message_id"] for event in events if event["event_type"] == "bounced"}),
               "unknown": sum(message["status"] == "unknown" for message in messages),
               "clicked_candidate": len({event["message_id"] for event in events if event["event_type"] == "clicked_candidate"}),
               "confirmed_human_interaction": len({event["message_id"] for event in events if event["classification"] == "confirmed_human"}),
               "reported": len({event["message_id"] for event in events if event["event_type"] == "reported"}),
               "definition_version": "awareness-metrics-v1", "denominator": "unique planned messages", "sandbox": True}
    return {**row, "messages": [{key: value for key, value in message.items() if key != "tracking_hash"} for message in messages],
            "events": events, "metrics": metrics, "notice": "本機模擬，未寄送真實郵件；供應商接受不代表已送達。"}


@router.get("/customer/recipient-groups")
def groups(ctx: Context = Depends(context), page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100)):
    ctx.require("campaign_manager")
    rows = all_rows(ctx.conn, m.recipient_groups, ctx.tenant_id, m.recipient_groups.c.owner_id == ctx.user_id)
    return paged([{**row, "recipient_count": len(all_rows(ctx.conn, m.recipients, ctx.tenant_id, m.recipients.c.group_id == row["id"]))} for row in rows], page, page_size)


@router.post("/customer/recipient-imports")
def import_recipients(payload: RecipientInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")

    def run():
        reader = csv.DictReader(StringIO(payload.csv.lstrip("\ufeff")))
        if not reader.fieldnames or not {"email", "department", "name"}.issubset(reader.fieldnames):
            fail(422, "RECIPIENT_COLUMNS", "CSV 欄位須包含 email,department,name。")
        if any(column.lower() in {"password", "passwd", "secret", "token"} for column in reader.fieldnames):
            fail(422, "SENSITIVE_FIELD_PROHIBITED", "名單不能包含密碼或憑證欄位。")
        valid, errors, seen = [], [], set()
        for index, row in enumerate(reader, 2):
            if index > 10002:
                fail(413, "RECIPIENT_LIMIT", "每次名單最多 10,000 列。")
            email = (row.get("email") or "").strip().lower()
            if not re.fullmatch(r"[^\s@,<>]+@[^\s@,<>]+\.[^\s@,<>]+", email) or len(email) > 254:
                errors.append({"row": index, "code": "invalid_email"})
            elif email in seen:
                errors.append({"row": index, "code": "duplicate_email"})
            else:
                seen.add(email)
                name, department = (row.get("name") or "").strip(), (row.get("department") or "").strip()
                if len(name) > 120 or len(department) > 120:
                    errors.append({"row": index, "code": "field_too_long"})
                else:
                    valid.append({"email": email, "name": name, "department": department})
        if not valid:
            return {"status": "invalid", "errors": errors, "valid_count": 0, "id": None}
        group = add(ctx.conn, m.recipient_groups, ctx.tenant_id, name=payload.name, owner_id=ctx.user_id)
        for recipient in valid:
            user = one(ctx.conn, m.users, None, m.users.c.email == recipient["email"])
            linked = user and one(ctx.conn, m.memberships, None, m.memberships.c.user_id == user["id"], m.memberships.c.tenant_id == ctx.tenant_id,
                                 m.memberships.c.role == "learner", m.memberships.c.active.is_(True))
            add(ctx.conn, m.recipients, ctx.tenant_id, group_id=group["id"], learner_id=user["id"] if linked else None, **recipient)
        ctx.audit("recipients.import", group["id"], f"valid={len(valid)}; invalid_or_duplicate={len(errors)}")
        return {**group, "status": "partial_success" if errors else "imported", "valid_count": len(valid), "errors": errors}
    return idempotent(ctx, request, "recipients.import", payload.model_dump(), run)


@router.get("/customer/campaigns")
def campaigns(ctx: Context = Depends(context), page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100)):
    ctx.require("campaign_manager")
    return paged([campaign_view(ctx, row) for row in all_rows(ctx.conn, m.campaigns, ctx.tenant_id, m.campaigns.c.owner_id == ctx.user_id)], page, page_size)


@router.post("/customer/recipient-imports/xlsx")
def import_recipient_workbook(payload: RecipientWorkbookInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    from defusedxml import ElementTree
    from openpyxl import load_workbook
    try:
        data = base64.b64decode(payload.content_base64, validate=True)
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(item.file_size for item in entries) > 8_000_000:
                raise ValueError("WORKBOOK_LIMIT")
            for item in entries:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in item.filename or ":" in item.filename or item.flag_bits & 1:
                    raise ValueError("UNSAFE_ARCHIVE")
                if item.file_size > 1_000_000 or item.file_size > max(item.compress_size, 1)*100:
                    raise ValueError("WORKBOOK_LIMIT")
                if "externallinks" in item.filename.lower() or item.filename.lower().endswith((".bin", ".vba")):
                    raise ValueError("ACTIVE_CONTENT")
                if item.filename.endswith((".xml", ".rels")):
                    root = ElementTree.fromstring(archive.read(item))
                    if any(element.attrib.get("TargetMode") == "External" for element in root.iter()):
                        raise ValueError("EXTERNAL_RELATIONSHIP")
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise ValueError("ONE_WORKSHEET_REQUIRED")
            sheet = book.worksheets[0]
            if sheet.max_row > 10001 or sheet.max_column > 10:
                raise ValueError("WORKBOOK_LIMIT")
            output = StringIO()
            writer = csv.writer(output)
            for index, cells in enumerate(sheet.iter_rows(), 1):
                if index > 10001 or any(cell.data_type == "f" for cell in cells):
                    raise ValueError("FORMULA_OR_ROW_LIMIT")
                writer.writerow(["" if cell.value is None else str(cell.value) for cell in cells])
        finally:
            book.close()
    except Exception:
        fail(422, "WORKBOOK_REJECTED", "僅接受單工作表、無公式/巨集/外部連結的 XLSX；請確認容量與欄位。")
    if len(output.getvalue()) > 2_000_000:
        fail(413, "RECIPIENT_LIMIT", "名單解壓縮後超出文字容量上限。")
    return import_recipients(RecipientInput(name=payload.name, csv=output.getvalue()), request, ctx)


@router.post("/customer/campaigns")
def create_campaign(payload: CampaignInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    group = owned(ctx, m.recipient_groups, payload.group_id)
    if group["owner_id"] != ctx.user_id:
        fail(404, "GROUP_SCOPE_DENIED", "受測群組未授權。")
    if aware(payload.scheduled_at) < m.now()-timedelta(minutes=1):
        fail(422, "SCHEDULE_IN_PAST", "排程時間不能早於現在。")
    if payload.remediation_course_id:
        ctx.require("training_manager")
        course_row(ctx.conn, payload.remediation_course_id)
    return idempotent(ctx, request, "campaign.create", payload.model_dump(),
                      lambda: add(ctx.conn, m.campaigns, ctx.tenant_id, owner_id=ctx.user_id, **payload.model_dump()))


@router.get("/customer/campaigns/{campaign_id}")
def campaign_detail(campaign_id: str, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    return campaign_view(ctx, owned_campaign(ctx, campaign_id))


@router.post("/customer/campaigns/{campaign_id}/schedule")
def schedule_campaign(campaign_id: str, payload: ConfirmationInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    row = owned_campaign(ctx, campaign_id)

    def run():
        if row["status"] != "draft":
            fail(409, "CAMPAIGN_STATE", "只有草稿活動可確認排程。")
        recipients = all_rows(ctx.conn, m.recipients, ctx.tenant_id, m.recipients.c.group_id == row["group_id"])
        if not recipients:
            fail(422, "EMPTY_RECIPIENTS", "需先加入有效受測名單。")
        expiry = aware(row["scheduled_at"])+timedelta(hours=1)
        if expiry <= m.now():
            fail(409, "CAMPAIGN_WINDOW_EXPIRED", "排程窗口已過，請建立新活動版本。")
        for recipient in recipients:
            message_id = m.uid()
            reservation = wallet.reserve(ctx.conn, ctx.tenant_id, "phishing", message_id, 1, expiry)
            tracking_token = m.uid()
            add(ctx.conn, m.message_plans, ctx.tenant_id, id=message_id, campaign_id=campaign_id,
                recipient_id=recipient["id"], reservation_id=reservation["id"], tracking_hash=digest(tracking_token))
            add(ctx.conn, m.jobs, id=tracking_token, tenant_id=ctx.tenant_id, kind="tracking_lookup", resource_id=message_id, status="active")
            add(ctx.conn, m.jobs, tenant_id=ctx.tenant_id, kind="sandbox_mail", resource_id=message_id, available_at=row["scheduled_at"])
        change(ctx.conn, m.campaigns, ctx.tenant_id, campaign_id, status="scheduled")
        ctx.audit("campaign.schedule", campaign_id, f"messages={len(recipients)}; sandbox authorization")
        return campaign_view(ctx, owned_campaign(ctx, campaign_id))
    return idempotent(ctx, request, f"campaign.schedule:{campaign_id}", payload.model_dump(), run)


def stop_campaign(ctx, row, action):
    if action == "resume":
        if row["status"] != "paused":
            fail(409, "CAMPAIGN_STATE", "只有暫停活動可繼續。")
        remaining = 0
        for message in all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.campaign_id == row["id"], m.message_plans.c.status == "planned"):
            reservation = owned(ctx, m.reservations, message["reservation_id"])
            expired = aware(reservation["expires_at"]) <= m.now()
            if expired:
                wallet.release(ctx.conn, ctx.tenant_id, reservation["id"], "campaign resumed after reservation expiry")
                change(ctx.conn, m.message_plans, ctx.tenant_id, message["id"], status="cancelled")
            else:
                remaining += 1
            for job in all_rows(ctx.conn, m.jobs, None, m.jobs.c.tenant_id == ctx.tenant_id, m.jobs.c.kind == "sandbox_mail", m.jobs.c.resource_id == message["id"], m.jobs.c.status == "paused"):
                change(ctx.conn, m.jobs, ctx.tenant_id, job["id"], status="cancelled" if expired else "queued", available_at=max(aware(row["scheduled_at"]), m.now()))
        change(ctx.conn, m.campaigns, ctx.tenant_id, row["id"], status="scheduled" if remaining else "observing")
    elif action == "pause":
        if row["status"] not in {"scheduled", "observing"}:
            fail(409, "CAMPAIGN_STATE", "活動目前無法暫停。")
        change(ctx.conn, m.campaigns, ctx.tenant_id, row["id"], status="paused")
        message_ids = [message["id"] for message in all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.campaign_id == row["id"])]
        for job in all_rows(ctx.conn, m.jobs, None, m.jobs.c.tenant_id == ctx.tenant_id, m.jobs.c.kind == "sandbox_mail", m.jobs.c.resource_id.in_(message_ids), m.jobs.c.status == "queued"):
            change(ctx.conn, m.jobs, ctx.tenant_id, job["id"], status="paused")
    else:
        if row["status"] == "cancelled":
            return campaign_view(ctx, row)
        for message in all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.campaign_id == row["id"]):
            if message["status"] == "planned":
                wallet.release(ctx.conn, ctx.tenant_id, message["reservation_id"], "campaign cancelled before acceptance")
                change(ctx.conn, m.message_plans, ctx.tenant_id, message["id"], status="cancelled")
                for job in all_rows(ctx.conn, m.jobs, None, m.jobs.c.tenant_id == ctx.tenant_id, m.jobs.c.resource_id == message["id"], m.jobs.c.kind == "sandbox_mail"):
                    change(ctx.conn, m.jobs, ctx.tenant_id, job["id"], status="cancelled")
        # Unknown acceptance stays reserved for reconciliation; accepted mail cannot be recalled.
        change(ctx.conn, m.campaigns, ctx.tenant_id, row["id"], status="cancelled")
    ctx.audit(f"campaign.{action}", row["id"])
    return campaign_view(ctx, owned_campaign(ctx, row["id"]))


@router.post("/customer/campaigns/{campaign_id}/cancel")
@router.post("/customer/campaigns/{campaign_id}/pause")
@router.post("/customer/campaigns/{campaign_id}/resume")
def campaign_action(campaign_id: str, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    row = owned_campaign(ctx, campaign_id)
    action = request.url.path.rsplit("/", 1)[-1]
    return idempotent(ctx, request, f"campaign.{action}:{campaign_id}", {}, lambda: stop_campaign(ctx, row, action))


@router.post("/development/campaigns/{campaign_id}/dispatch")
def dispatch(campaign_id: str, request: Request, payload: DispatchInput = DispatchInput(), ctx: Context = Depends(context)):
    if settings().app_env not in {"development", "test"}:
        fail(404, "NOT_FOUND", "找不到此操作。")
    ctx.require("campaign_manager")
    row = owned_campaign(ctx, campaign_id)

    def run():
        if row["status"] not in {"scheduled", "observing"}:
            fail(409, "CAMPAIGN_STATE", "活動需先確認排程。")
        from .worker import settle_message
        for message in all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.campaign_id == campaign_id,
                                 m.message_plans.c.status == "planned"):
            settle_message(ctx.conn, ctx.tenant_id, message["id"], payload.outcome)
        change(ctx.conn, m.campaigns, ctx.tenant_id, campaign_id, status="observing")
        ctx.audit("campaign.simulate", campaign_id, f"outcome={payload.outcome}; real_mail=0")
        return campaign_view(ctx, owned_campaign(ctx, campaign_id))
    return idempotent(ctx, request, f"campaign.simulate:{campaign_id}", payload.model_dump(), run)


@router.post("/customer/suspicious-mail-reports")
def suspicious_report(payload: SuspiciousInput, request: Request, ctx: Context = Depends(context)):
    if payload.campaign_id:
        ctx.require("campaign_manager")
        owned_campaign(ctx, payload.campaign_id)
    return idempotent(ctx, request, "suspicious.report", payload.model_dump(),
                      lambda: add(ctx.conn, m.tickets, ctx.tenant_id, subject="可疑郵件回報", text=payload.description, actor_id=ctx.user_id))


@router.post("/customer/campaigns/{campaign_id}/remediation")
def remediation(campaign_id: str, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    ctx.require("training_manager")
    row = owned_campaign(ctx, campaign_id)
    if not row["remediation_course_id"]:
        fail(409, "REMEDIATION_NOT_CONFIGURED", "活動尚未選擇補救課程。")

    def run():
        events = all_rows(ctx.conn, m.raw_events, ctx.tenant_id, m.raw_events.c.campaign_id == campaign_id,
                          m.raw_events.c.classification == "confirmed_human")
        enrolled, skipped = [], 0
        for message_id in sorted({event["message_id"] for event in events}):
            message = owned(ctx, m.message_plans, message_id)
            recipient = owned(ctx, m.recipients, message["recipient_id"])
            if not recipient["learner_id"]:
                skipped += 1
                continue
            enrolled.append(assign(ctx, row["remediation_course_id"], recipient["learner_id"], f"remediation:{campaign_id}"))
        return {"items": enrolled, "unlinked_recipients": skipped, "trigger": "confirmed_human_only"}
    return idempotent(ctx, request, f"campaign.remediation:{campaign_id}", {}, run)


@router.post("/customer/campaigns/{campaign_id}/events/{event_id}/review")
def review_event(campaign_id: str, event_id: str, payload: EventReviewInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("campaign_manager")
    owned_campaign(ctx, campaign_id)
    event = owned(ctx, m.raw_events, event_id)
    if event["campaign_id"] != campaign_id:
        fail(404, "EVENT_SCOPE_DENIED", "事件不屬於此活動。")
    if event["event_type"] not in {"clicked_candidate", "human_interaction"}:
        fail(409, "EVENT_TYPE_NOT_REVIEWABLE", "只有互動候選事件可判讀為人類或自動設備。")

    def run():
        change(ctx.conn, m.raw_events, ctx.tenant_id, event_id, classification=payload.classification)
        ctx.audit("campaign.event_review", event_id, f"{event['classification']} -> {payload.classification}; evidence={payload.reason}")
        return owned(ctx, m.raw_events, event_id)
    return idempotent(ctx, request, f"campaign.event_review:{event_id}", payload.model_dump(), run)


@router.get("/internal/phishing/messages")
def messages_for_reconciliation(ctx: Context = Depends(context), page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=100)):
    ctx.require("finance")
    return paged([{key: value for key, value in row.items() if key != "tracking_hash"} for row in
                  all_rows(ctx.conn, m.message_plans, ctx.tenant_id, m.message_plans.c.status == "unknown")], page, page_size)


@router.post("/internal/phishing/messages/{message_id}/reconcile")
def reconcile_delivery(message_id: str, payload: ReconciliationInput, request: Request, ctx: Context = Depends(context)):
    ctx.require("finance")
    message = owned(ctx, m.message_plans, message_id)
    if payload.occurred_at < aware(message["created_at"]):
        fail(422, "INVALID_PROVIDER_TIME", "供應商時間不能早於此寄送規劃。")

    def run():
        if message["status"] != "unknown":
            fail(409, "MESSAGE_NOT_UNKNOWN", "只有供應商接受狀態未知的項目需要核對。")
        from .worker import reconcile_message
        return reconcile_message(ctx.conn, ctx.tenant_id, message_id, payload.outcome, payload.provider_reference,
                                 payload.occurred_at, payload.reason, ctx.user_id)
    return idempotent(ctx, request, f"mail.reconcile:{message_id}", payload.model_dump(), run)
