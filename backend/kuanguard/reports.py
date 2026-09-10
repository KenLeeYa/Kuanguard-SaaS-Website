"""One approved data snapshot, multiple auditable development report artifacts.

No Office process, network service, source repository, or arbitrary template runs
inside this module. Production template qualification is a separate release gate.
"""

from __future__ import annotations

import copy
import csv
import hashlib
import json
import os
from collections import Counter
from pathlib import Path
import re
import shutil
import tempfile
from xml.sax.saxutils import escape

from .parsers import SEVERITIES

REPORT_VERSION = "kuanguard-reports/1.0.0"
RISK_LEVELS = SEVERITIES[:4]
LEVEL_ZH = {"Critical": "嚴重", "High": "高", "Medium": "中", "Low": "低", "Informational": "資訊", "Unknown": "未知"}
STATUS_ZH = {"completed": "已完成", "not_tested": "未檢測", "failed": "失敗", "insufficient": "證據不足",
             "passed": "通過", "not_applicable": "不適用"}
RETEST_ZH = {"new": "新增", "persistent": "仍存在", "verified_remediated": "已驗證修復", "reappeared": "再次出現",
             "not_covered": "未覆蓋", "indeterminate": "無法判定"}
COLORS = {"Critical": "B91C1C", "High": "C26011", "Medium": "1D4ED8", "Low": "166534", "Informational": "475569", "Unknown": "64748B"}
ARTIFACT_NAMES = ("report.docx", "report.pdf", "summary.pptx", "summary.xlsx", "all.csv")
DISCLOSURE = "合成示範資料，僅供本機開發驗證；未使用正式核准模板，不代表實際客戶檢測。"


def _canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _hash(value):
    return hashlib.sha256(_canonical(value)).hexdigest()


def _risk_findings(findings, service):
    return [f for f in findings if f["severity"] in RISK_LEVELS and (service != "SHC" or f.get("check_status") == "failed")]


def _merge_sources(parsed, context):
    sources = [parsed, *context.get("additional_sources", [])]
    if any(source.get("errors") for source in sources):
        raise ValueError("不可從包含匯入錯誤的資料建立快照")
    if len(sources) == 1:
        return copy.deepcopy(parsed), [{"sha256": parsed["source_hash"], "format": parsed.get("metadata", {}).get("format", "unknown")}], []
    precedence = context.get("source_precedence", "csv")
    if precedence not in {"csv", "first", "merge"}:
        raise ValueError("source_precedence 必須是 csv、first 或 merge")
    if precedence == "csv" and not any(source.get("metadata", {}).get("format") == "nessus-csv" for source in sources):
        precedence = "merge"
    chosen = next((s for s in sources if s.get("metadata", {}).get("format") == "nessus-csv"), parsed) if precedence == "csv" else parsed
    if precedence == "merge":
        chosen = copy.deepcopy(parsed)
        merged = {}
        for source in sources:
            for finding in source["findings"]:
                merged.setdefault(finding["fingerprint"], copy.deepcopy(finding))
        chosen["findings"] = list(merged.values())
    comparison = []
    chosen_keys = {f["fingerprint"] for f in chosen["findings"]}
    for source in sources:
        other = {f["fingerprint"] for f in source["findings"]}
        selected_by_id = {f["fingerprint"]: f for f in chosen["findings"]}
        changes = [{"fingerprint": f["fingerprint"], "fields": [key for key in ("severity", "title", "description", "solution")
                    if f.get(key) != selected_by_id[f["fingerprint"]].get(key)]} for f in source["findings"] if f["fingerprint"] in selected_by_id]
        changes = [change for change in changes if change["fields"]]
        comparison.append({"source_hash": source["source_hash"], "selected": source is chosen or precedence == "merge",
                           "only_in_selected": len(chosen_keys - other), "only_in_source": len(other - chosen_keys),
                           "changed_occurrences": changes,
                           "reason": {"csv": "專案明定 CSV 優先", "first": "專案明定第一來源優先", "merge": "核定來源合併，重複實例保留第一來源"}[precedence]})
    result = copy.deepcopy(chosen)
    # Keep independently proven scope, while conflicting failure wins over completion.
    coverage = {}
    order = {"not_tested": 0, "insufficient": 1, "completed": 2, "failed": 3}
    for source in sources:
        for row in source["coverage"]:
            if order[row["status"]] >= order.get(coverage.get(row["asset"]), -1):
                coverage[row["asset"]] = row["status"]
    result["coverage"] = [{"asset": a, "status": s} for a, s in sorted(coverage.items())]
    return result, [{"sha256": s["source_hash"], "format": s.get("metadata", {}).get("format", "unknown")} for s in sources], comparison


def _occurrence_key(finding, context):
    assets = context.get("stable_assets", {})
    paths = context.get("path_aliases", {})
    asset = assets.get(finding["asset"], finding["asset"])
    if finding.get("source_fingerprint"):
        parts = [finding.get("method"), finding["source_id"], finding["source_fingerprint"]]
    else:
        parts = [paths.get(asset, asset), finding["source_id"], finding.get("port"), finding.get("protocol"),
                 finding.get("location_hash"), finding.get("http_method"), finding.get("parameter"), finding.get("method"),
                 finding.get("line"), finding.get("column")]
    return _hash(parts)


def _retest(findings, coverage, context, metadata):
    baseline = context.get("baseline_snapshot")
    if not baseline:
        return []
    current = {f["occurrence_id"]: f for f in findings}
    original = {_occurrence_key(f, context): f for f in baseline["findings"]}
    coverage_map = {r["asset"]: r["status"] for r in coverage}
    verified = {r.get("occurrence_id"): r for r in context.get("verified_absent", [])}
    disabled = set(metadata.get("disabled_rules", []))
    excluded = metadata.get("excluded_paths", [])
    tested_rules = set(context.get("tested_rule_ids", []))
    rows = []
    for identity, finding in current.items():
        old = original.get(identity)
        state = "reappeared" if old and old.get("retest_status") == "verified_remediated" else "persistent" if old else "new"
        rows.append({"occurrence_id": identity, "asset": finding["asset"], "source_id": finding["source_id"],
                     "status": state, "evidence": "當次來源仍有此發現", "baseline_snapshot_id": baseline.get("snapshot_id")})
    for identity, finding in original.items():
        if identity in current:
            continue
        state, reason = "indeterminate", "未提供充分且經覆核的個別修復驗證"
        proof = verified.get(identity, {})
        if coverage_map.get(finding["asset"]) != "completed":
            state, reason = "not_covered", "原資產範圍未完整完成"
        elif finding.get("rule_id") in disabled or any(finding["asset"].startswith(p) for p in excluded):
            state, reason = "not_covered", "原規則停用或原路徑排除"
        elif finding.get("method", "").startswith("sarif:") and finding.get("rule_id") not in tested_rules:
            state, reason = "not_covered", "缺少原規則本次執行證據"
        elif proof.get("evidence") and proof.get("reviewer_id") and proof.get("method") == finding.get("method") and proof.get("scope_confirmed") is True:
            state, reason = "verified_remediated", str(proof["evidence"])
        rows.append({"occurrence_id": identity, "asset": finding["asset"], "source_id": finding["source_id"], "status": state,
                     "evidence": reason, "baseline_snapshot_id": baseline.get("snapshot_id")})
    return sorted(rows, key=lambda row: (row["asset"], row["source_id"]))


def build_snapshot(parsed: dict, context: dict) -> dict:
    """Create a detached, hash-addressed normalized payload; never alter inputs."""
    selected, sources, comparison = _merge_sources(parsed, context)
    service = context.get("service_code", selected.get("metadata", {}).get("service_code", "VA"))
    findings = copy.deepcopy(selected["findings"])
    if service == "PT" and any(f.get("reviewed") is not True for f in findings):
        raise ValueError("PT 發現尚未完成人工覆核，禁止產生核定快照")
    for finding in findings:
        if finding.get("severity") not in SEVERITIES:
            raise ValueError("快照嚴重度不符合 schema")
        finding["occurrence_id"] = _occurrence_key(finding, context)
    findings.sort(key=lambda f: (SEVERITIES.index(f["severity"]), f["asset"], f["source_id"], str(f.get("port", "")), f.get("location", "")))
    coverage = copy.deepcopy(selected["coverage"])
    risk = _risk_findings(findings, service)
    counts = {level: sum(f["severity"] == level for f in risk) for level in RISK_LEVELS}
    level_counts = {level: sum(f["severity"] == level for f in findings) for level in SEVERITIES}
    planned = selected.get("metadata", {}).get("planned_assets", [r["asset"] for r in coverage])
    planned = sorted(set(planned))
    completed = sum(row["status"] == "completed" and row["asset"] in planned for row in coverage)
    hosts = []
    for row in coverage:
        local = [f for f in risk if f["asset"] == row["asset"]]
        hosts.append({**row, "planned": row["asset"] in planned, "risk_count": len(local),
                      "severity_counts": {level: sum(f["severity"] == level for f in local) for level in RISK_LEVELS}})
    hosts.sort(key=lambda row: (-row["risk_count"], row["asset"]))
    statistics = {"severity_counts": counts, "source_severity_counts": level_counts, "occurrence_count": len(risk),
                  "definition_count": len({(f.get("method"), f["source_id"]) for f in risk}),
                  "affected_asset_count": len({f["asset"] for f in risk}), "source_record_count": len(findings),
                  "informational_count": level_counts["Informational"], "unknown_count": level_counts["Unknown"],
                  "planned_asset_count": len(planned), "completed_asset_count": completed,
                  "coverage_ratio": completed / len(planned) if planned else None,
                  "coverage_counts": dict(Counter(row["status"] for row in coverage)), "assets": hosts,
                  "denominator": "預定範圍資產數；未預定的新資產另列，不擴大分母", "risk_denominator": "四級弱點實例，資訊與未知另外列示"}
    if service == "SHC":
        checks = Counter(f.get("check_status", "insufficient") for f in findings)
        denominator = len(findings) - checks["not_applicable"]
        statistics["shc"] = {"check_counts": {key: checks[key] for key in ("passed", "failed", "not_applicable", "not_tested", "insufficient")},
                             "applicable_count": denominator, "pass_rate": checks["passed"] / denominator if denominator else None,
                             "evidence_coverage_rate": (checks["passed"] + checks["failed"]) / denominator if denominator else None,
                             "denominator": "全部適用檢核，包含未檢測及證據不足；不適用排除分母"}
    snapshot = {"schema_version": "kuanguard.snapshot/1", "generator_version": REPORT_VERSION,
                "parser_version": selected["parser_version"], "normalized_dataset_version": "1.0", "sources": sources,
                "source_comparison": comparison, "source_hash": selected["source_hash"], "service_code": service,
                "batch_id": str(context.get("batch_id", "unassigned")), "scope_version": str(context.get("scope_version", "1")),
                "report_version": str(context.get("report_version", "1")), "reviewer_id": str(context.get("reviewer_id", "")),
                "reviewed_at": str(context.get("reviewed_at", "")), "source_timezone": str(context.get("source_timezone", "UTC")),
                "business_timezone": "Asia/Taipei", "project_name": str(context.get("project_name", "KUANGUARD 檢測示範")),
                "trace_id": str(context.get("trace_id", context.get("batch_id", "unassigned"))),
                "synthetic": context.get("synthetic") is True,
                "template_version": "kuanguard-development/1", "rule_version": str(context.get("rule_version", "no-exclusions/1")),
                "translation_version": "original-source/no-ai", "knowledge_version": "source-only/1", "findings": findings,
                "coverage": coverage, "planned_assets": planned, "statistics": statistics,
                "metadata": selected.get("metadata", {}), "warnings": selected.get("warnings", []),
                "retest": _retest(findings, coverage, context, selected.get("metadata", {}))}
    snapshot["publication_gates"] = {"review_approved": bool(snapshot["reviewer_id"] and snapshot["reviewed_at"]),
                                     "production_ready": False, "production_blockers": ["尚未提供核准模板與使用權紀錄", "Office 全頁渲染與人工版面覆核待完成"]}
    snapshot["snapshot_hash"] = _hash(snapshot)
    snapshot["snapshot_id"] = "snap_" + snapshot["snapshot_hash"][:24]
    return snapshot


def _verify_snapshot(snapshot):
    copy_data = {key: value for key, value in snapshot.items() if key not in {"snapshot_hash", "snapshot_id"}}
    if _hash(copy_data) != snapshot.get("snapshot_hash"):
        raise ValueError("SNAPSHOT_HASH_MISMATCH: 快照建立後已被更動")
    if not snapshot.get("publication_gates", {}).get("review_approved"):
        raise ValueError("REVIEW_REQUIRED: 必須有覆核者及覆核時間")
    if not snapshot.get("synthetic"):
        raise ValueError("PRODUCTION_TEMPLATE_NOT_QUALIFIED: 正式模板與全頁 QA 尚未通過")


def _safe_cell(value):
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def _csv_rows(snapshot):
    headers = ["snapshot_id", "report_version", "service", "occurrence_id", "source_id", "asset", "severity", "title", "description",
               "solution", "port", "location", "check_status", "evidence", "source_hash", "synthetic", "disclosure", "rule_id", "line", "column", "source_fingerprint"]
    rows = [[snapshot["snapshot_id"], snapshot["report_version"], snapshot["service_code"], f["occurrence_id"], f["source_id"], f["asset"],
             f["severity"], f["title"], f["description"], f["solution"], f.get("port"), f.get("location"), f.get("check_status"),
             f.get("evidence"), snapshot["source_hash"], snapshot["synthetic"], DISCLOSURE if snapshot["synthetic"] else "",
             f.get("rule_id"), f.get("line"), f.get("column"), f.get("source_fingerprint")] for f in snapshot["findings"]]
    return headers, rows


def _write_csv(snapshot, path):
    headers, rows = _csv_rows(snapshot)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows([_safe_cell(value) for value in row] for row in rows)


def _write_xlsx(snapshot, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["KUANGUARD 檢測摘要", DISCLOSURE])
    summary.append(["snapshot_id", snapshot["snapshot_id"]])
    summary.append(["snapshot_hash", snapshot["snapshot_hash"]])
    summary.append(["report_version", snapshot["report_version"]])
    summary.append(["嚴重度", "弱點實例數"])
    for severity in RISK_LEVELS:
        summary.append([f"{LEVEL_ZH[severity]} {severity}", snapshot["statistics"]["severity_counts"][severity]])
        summary.cell(summary.max_row, 1).font = Font(color=COLORS[severity])
    for label, key in [("弱點種類", "definition_count"), ("弱點實例", "occurrence_count"), ("受影響資產", "affected_asset_count"),
                       ("Informational 不納入四級", "informational_count"), ("Unknown 待釐清", "unknown_count"),
                       ("預定資產", "planned_asset_count"), ("完成資產", "completed_asset_count")]:
        summary.append([label, snapshot["statistics"][key]])
    summary.append(["統計口徑", snapshot["statistics"]["risk_denominator"]])
    data = book.create_sheet("Findings")
    headers, rows = _csv_rows(snapshot)
    data.append(headers)
    for row in rows:
        data.append([_safe_cell(value) for value in row])
    scope = book.create_sheet("Coverage")
    scope.append(["asset", "status", "planned", "Critical", "High", "Medium", "Low", "total"])
    for row in snapshot["statistics"]["assets"]:
        scope.append([row["asset"], row["status"], row["planned"], *[row["severity_counts"][level] for level in RISK_LEVELS], row["risk_count"]])
    scope.append(["TOTAL", "", "", *[snapshot["statistics"]["severity_counts"][level] for level in RISK_LEVELS], snapshot["statistics"]["occurrence_count"]])
    retest = book.create_sheet("Retest")
    retest.append(["occurrence_id", "asset", "source_id", "status", "evidence", "baseline_snapshot_id"])
    for row in snapshot["retest"]:
        retest.append([_safe_cell(row.get(key)) for key in [cell.value for cell in retest[1]]])
    if "shc" in snapshot["statistics"]:
        sheet = book.create_sheet("SHC")
        sheet.append(["status", "count"])
        for status, count in snapshot["statistics"]["shc"]["check_counts"].items():
            sheet.append([status, count])
        sheet.append(["denominator", snapshot["statistics"]["shc"]["denominator"]])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="0B2545")
            cell.font = Font(color="FFFFFF", bold=True)
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 26 if index < 8 else 48
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.value = _safe_cell(cell.value)
        sheet.sheet_view.zoomScale = 85
    summary.column_dimensions["B"].width = 85
    book.properties.title = "KUANGUARD 檢測摘要"
    book.properties.subject = snapshot["snapshot_id"]
    book.save(path)


def _set_word_font(run, size=14, bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    fonts.set(qn("w:eastAsia"), "標楷體")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def _word_table(doc, headers, rows):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, headers):
        cell.text = str(heading)
        for run in cell.paragraphs[0].runs:
            _set_word_font(run, 11, True)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value) if value is not None else "未提供"
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_word_font(run, 11)
    return table


def _write_docx(snapshot, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(14)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    for heading in ("Title", "Heading 1", "Heading 2"):
        doc.styles[heading].font.color.rgb = RGBColor(0, 0, 0)
    def paragraph(text, style=None, size=14):
        p = doc.add_paragraph(style=style)
        _set_word_font(p.add_run(str(text)), size)
        return p
    paragraph(f"KUANGUARD {snapshot['service_code']} 檢測報告", "Title", 24)
    paragraph(snapshot["project_name"], size=18)
    paragraph(DISCLOSURE, size=12)
    paragraph(f"版本 {snapshot['report_version']}　覆核 {snapshot['reviewer_id']}　時間 {snapshot['reviewed_at']}", size=11)
    paragraph(f"快照 {snapshot['snapshot_id']}", size=10)
    paragraph("本報告保留檢測範圍、證據限制及改善資訊。資訊級與未知風險分別列示，未檢測與證據不足不視為安全。")
    paragraph("目錄", "Heading 1", 18)
    toc = paragraph("", size=12)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    toc._p.append(field)
    paragraph("目錄欄位需經正式 Office 渲染更新；本示範以下提供完整章節。", size=10)
    doc.add_page_break()
    paragraph("檢測範圍與統計", "Heading 1", 18)
    _word_table(doc, ["資產", "檢測狀態", "四級弱點數"], [[r["asset"], STATUS_ZH[r["status"]], r["risk_count"]] for r in snapshot["statistics"]["assets"]])
    paragraph(f"預定資產 {snapshot['statistics']['planned_asset_count']}；完成 {snapshot['statistics']['completed_asset_count']}。", size=12)
    _word_table(doc, ["嚴重度", "弱點實例數"], [[LEVEL_ZH[level], snapshot["statistics"]["severity_counts"][level]] for level in RISK_LEVELS])
    paragraph(f"弱點種類 {snapshot['statistics']['definition_count']}；弱點實例 {snapshot['statistics']['occurrence_count']}；受影響資產 {snapshot['statistics']['affected_asset_count']}。", size=12)
    paragraph(f"Informational {snapshot['statistics']['informational_count']}；Unknown {snapshot['statistics']['unknown_count']}，均不納入四級風險。", size=12)
    if "shc" in snapshot["statistics"]:
        paragraph("健診符合狀態", "Heading 1", 18)
        _word_table(doc, ["狀態", "檢核數"], [[STATUS_ZH[k], v] for k, v in snapshot["statistics"]["shc"]["check_counts"].items()])
        paragraph(snapshot["statistics"]["shc"]["denominator"], size=12)
    if snapshot["service_code"] == "PT":
        paragraph("授權與測試方法", "Heading 1", 18)
        for label, key in [("授權 reference", "authorization_reference"), ("方法", "methodology"), ("時段", "test_window"), ("停止條件", "stop_conditions")]:
            paragraph(f"{label}：{snapshot['metadata'].get(key, '未提供')}", size=12)
    if snapshot["service_code"] == "SOURCE":
        paragraph("源碼檢測範圍", "Heading 1", 18)
        paragraph("本批次僅匯入 SARIF 結果。來源程式碼未在本服務執行。", size=12)
        paragraph(f"工具 {json.dumps(snapshot['metadata'].get('tools', []), ensure_ascii=False)}", size=12)
        paragraph(f"停用規則 {', '.join(snapshot['metadata'].get('disabled_rules', [])) or '來源未列示'}", size=12)
        paragraph(f"排除路徑 {', '.join(snapshot['metadata'].get('excluded_paths', [])) or '來源未列示'}", size=12)
    paragraph("前十大風險實例", "Heading 1", 18)
    top = _risk_findings(snapshot["findings"], snapshot["service_code"])[:10]
    _word_table(doc, ["風險", "來源編號", "資產與名稱"], [[LEVEL_ZH[f["severity"]], f["source_id"], f["asset"] + "\n" + f["title"]] for f in top] or [["N/A", "N/A", "本次沒有四級風險實例；仍須查看範圍限制。"]])
    paragraph("檢測明細與改善", "Heading 1", 18)
    for level in SEVERITIES:
        paragraph(f"{LEVEL_ZH[level]}級結果", "Heading 2", 16)
        group = [f for f in snapshot["findings"] if f["severity"] == level]
        if not group:
            paragraph("N/A", size=12)
        for finding in group:
            rows = [["來源與名稱", finding["source_id"] + "　" + finding["title"]], ["資產", finding["asset"]],
                    ["位置與連接埠", (finding.get("location") or "未提供") + " / " + str(finding.get("port") if finding.get("port") is not None else "未提供")],
                    ["說明與影響", finding["description"] or "來源未提供"], ["改善建議", finding["solution"] or "來源未提供"],
                    ["證據", finding.get("evidence") or "來源未提供"]]
            for label, key in [("健診狀態", "check_status"), ("預期值", "expected"), ("觀測值", "observed"),
                               ("前置條件", "preconditions"), ("重現步驟", "reproduction"), ("風險理由", "risk_reason"),
                               ("來源行號", "line"), ("規則編號", "rule_id")]:
                if key in finding:
                    rows.append([label, finding.get(key) or "未提供"])
            _word_table(doc, ["欄位", "結果"], rows)
            paragraph("", size=4)
    paragraph("複測與資料追溯", "Heading 1", 18)
    if snapshot["retest"]:
        _word_table(doc, ["來源", "資產", "複測狀態", "依據"], [[r["source_id"], r["asset"], RETEST_ZH[r["status"]], r["evidence"]] for r in snapshot["retest"]])
    else:
        paragraph("本次未提供初測關聯，未作修復結論。", size=12)
    for source in snapshot["sources"]:
        paragraph(f"來源 SHA256 {source['sha256']}", size=9)
    paragraph(f"快照 SHA256 {snapshot['snapshot_hash']}", size=9)
    paragraph(f"Parser {snapshot['parser_version']}；模板 {snapshot['template_version']}；規則 {snapshot['rule_version']}", size=10)
    header = section.header.paragraphs[0]
    _set_word_font(header.add_run("KUANGUARD　合成示範報告"), 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_word_font(footer.add_run(f"{snapshot['snapshot_id']}　頁 "), 8)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    settings = OxmlElement("w:updateFields")
    settings.set(qn("w:val"), "true")
    doc.settings.element.append(settings)
    doc.core_properties.title = f"KUANGUARD {snapshot['service_code']} 檢測報告"
    doc.core_properties.subject = snapshot["snapshot_id"]
    doc.core_properties.author = "KUANGUARD"
    doc.save(path)


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    root = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    chinese, latin = root / "kaiu.ttf", root / "times.ttf"
    if chinese.exists() and latin.exists():
        pdfmetrics.registerFont(TTFont("KGChinese", str(chinese)))
        pdfmetrics.registerFont(TTFont("KGLatin", str(latin)))
        return "KGChinese", "KGLatin", {"status": "available_local", "chinese": "標楷體", "latin": "Times New Roman", "licensing_for_container_distribution": "unverified"}
    pdfmetrics.registerFont(UnicodeCIDFont("MSung-Light"))
    return "MSung-Light", "Times-Roman", {"status": "fallback_not_template_equivalent", "chinese": "MSung-Light", "latin": "Times-Roman", "production_blocker": "核准字型缺少，示範替代字型不得宣稱等同模板"}


def _pdf_markup(value, chinese, latin):
    parts = re.split(r"([\x20-\x7e]+)", str(value))
    return "".join(f'<font name="{latin if re.fullmatch(r"[\x20-\x7e]+", p) else chinese}">{escape(p).replace(chr(10), "<br/>")}</font>' for p in parts if p)


def _write_pdf(snapshot, path):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    chinese, latin, font_status = _pdf_fonts()
    styles = {"body": ParagraphStyle("body", fontName=chinese, fontSize=14, leading=21, wordWrap="CJK", spaceAfter=8),
              "small": ParagraphStyle("small", fontName=chinese, fontSize=10, leading=15, wordWrap="CJK", spaceAfter=6),
              "heading": ParagraphStyle("heading", fontName=chinese, fontSize=19, leading=26, spaceBefore=14, spaceAfter=10, keepWithNext=True),
              "subheading": ParagraphStyle("subheading", fontName=chinese, fontSize=16, leading=22, spaceBefore=8, spaceAfter=6, keepWithNext=True),
              "title": ParagraphStyle("title", fontName=chinese, fontSize=27, leading=36, spaceAfter=16),
              "cell": ParagraphStyle("cell", fontName=chinese, fontSize=10, leading=15, wordWrap="CJK", alignment=TA_LEFT)}
    def p(text, style="body"):
        return Paragraph(_pdf_markup(text, chinese, latin), styles[style])
    story = []
    def table(headers, rows, widths):
        cells = [[p(value, "cell") for value in row] for row in [headers, *rows]]
        tab = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
        tab.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EFF5")), ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8DA2B6")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
        story.append(tab)
        story.append(Spacer(1, 10))
    story += [p(f"KUANGUARD {snapshot['service_code']}", "title"), p("檢測結果與改善報告", "title"), p(snapshot["project_name"]), p(DISCLOSURE, "small"),
              p(f"版本 {snapshot['report_version']}　覆核 {snapshot['reviewer_id']}", "small"), p(snapshot["reviewed_at"], "small"),
              p(snapshot["snapshot_id"], "small"), p("本次結論以範圍與證據為限。未檢測及證據不足不視為安全。"),
              p("文件導覽", "heading"), p("檢測範圍與統計\n檢測明細與改善\n複測狀態與資料追溯"), PageBreak(), p("檢測範圍與統計", "heading")]
    table(["資產", "狀態", "四級弱點數"], [[r["asset"], STATUS_ZH[r["status"]], r["risk_count"]] for r in snapshot["statistics"]["assets"]], [285, 105, 93])
    table(["嚴重度", "弱點實例數"], [[f"{LEVEL_ZH[level]} {level}", snapshot["statistics"]["severity_counts"][level]] for level in RISK_LEVELS], [320, 163])
    stats = snapshot["statistics"]
    story += [p(f"弱點種類 {stats['definition_count']}；弱點實例 {stats['occurrence_count']}；受影響資產 {stats['affected_asset_count']}。", "small"),
              p(f"預定資產 {stats['planned_asset_count']}；完成 {stats['completed_asset_count']}。資訊級 {stats['informational_count']}；未知 {stats['unknown_count']}。", "small"),
              p(stats["risk_denominator"], "small")]
    if "shc" in stats:
        story.append(p("健診符合狀態", "heading"))
        table(["判定", "檢核數"], [[STATUS_ZH[k], v] for k, v in stats["shc"]["check_counts"].items()], [320, 163])
        story.append(p(stats["shc"]["denominator"], "small"))
    if snapshot["service_code"] == "PT":
        story.append(p("授權與測試方法", "heading"))
        for key in ("authorization_reference", "methodology", "test_window", "stop_conditions"):
            story.append(p(f"{key}: {snapshot['metadata'].get(key, '未提供')}", "small"))
    if snapshot["service_code"] == "SOURCE":
        story += [p("源碼範圍限制", "heading"), p("僅匯入來源結果，未執行原始程式碼。規則停用或路徑排除時，缺少發現不代表修復。", "small")]
    story += [PageBreak(), p("檢測明細與改善", "heading")]
    for level in SEVERITIES:
        group = [f for f in snapshot["findings"] if f["severity"] == level]
        if not group:
            story.append(KeepTogether([p(f"{LEVEL_ZH[level]}級結果", "subheading"), p("N/A", "small")]))
        for index, finding in enumerate(group):
            endpoint = finding["asset"]
            if finding.get("port") is not None:
                endpoint = (f"[{endpoint}]" if ":" in endpoint and "://" not in endpoint else endpoint) + f":{finding['port']}"
            lead = [p(f"{LEVEL_ZH[level]}級結果", "subheading")] if index == 0 else []
            story.append(KeepTogether([*lead, p(f"{finding['source_id']} {finding['title']}"), p(endpoint + "　" + (finding.get("location") or ""), "small")]))
            # Long evidence is paragraph flow, never one unsplittable oversized table row.
            for label, key in [("說明", "description"), ("改善建議", "solution"), ("證據", "evidence"), ("健診狀態", "check_status"),
                               ("預期值", "expected"), ("觀測值", "observed"), ("前置條件", "preconditions"), ("重現步驟", "reproduction"), ("風險理由", "risk_reason"),
                               ("來源行號", "line"), ("規則編號", "rule_id")]:
                if key in finding:
                    story.append(p(label + "：" + (str(finding.get(key)) or "來源未提供"), "small"))
            story.append(Spacer(1, 8))
    trace = [p("複測與資料追溯", "heading")]
    for row in snapshot["retest"]:
        trace.append(p(f"{row['source_id']}　{row['asset']}　{RETEST_ZH[row['status']]}\n{row['evidence']}", "small"))
    if not snapshot["retest"]:
        trace.append(p("未提供初測關聯，未作修復結論。", "small"))
    trace += [p("快照 SHA256 " + snapshot["snapshot_hash"], "small"), p("來源 SHA256 " + snapshot["source_hash"], "small"),
              p("Parser " + snapshot["parser_version"] + "\n模板 " + snapshot["template_version"], "small")]
    story.append(KeepTogether(trace))
    def page_decoration(canvas, doc):
        canvas.setFont(chinese, 8)
        canvas.drawString(56, 815, "KUANGUARD 合成示範報告")
        canvas.setFont(latin, 8)
        canvas.drawString(56, 27, snapshot["snapshot_id"])
        canvas.drawRightString(539, 27, str(doc.page))
    document = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=56, leftMargin=56, topMargin=49, bottomMargin=45,
                                 title=f"KUANGUARD {snapshot['service_code']} Report", author="KUANGUARD", subject=snapshot["snapshot_id"])
    document.build(story, onFirstPage=page_decoration, onLaterPages=page_decoration)
    return font_status


def _text_chunks(text, capacity=330):
    text = str(text)
    chunks = []
    while text:
        if len(text) <= capacity:
            chunks.append(text)
            break
        cut = max(text.rfind("\n", 0, capacity), text.rfind("。", 0, capacity), text.rfind(" ", 0, capacity))
        if cut < capacity // 2:
            cut = capacity
        else:
            cut += 1
        chunks.append(text[:cut])
        text = text[cut:]
    return chunks or ["N/A"]


def _write_pptx(snapshot, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    def box(slide, text, left, top, width, height, size, color="152B43"):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        for index, line in enumerate(str(text).split("\n")):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.space_after = Pt(8)
            para.line_spacing = 1.5
            for segment in re.split(r"([\x20-\x7e]+)", line):
                if not segment:
                    continue
                run = para.add_run()
                run.text = segment
                run.font.name = "Times New Roman" if re.fullmatch(r"[\x20-\x7e]+", segment) else "Microsoft JhengHei"
                run.font.size = Pt(size)
                run.font.color.rgb = RGBColor.from_string(color)
                east = OxmlElement("a:ea")
                east.set("typeface", "Microsoft JhengHei")
                run._r.get_or_add_rPr().append(east)
        return shape
    def slide(title, body, body_size=24):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        box(page, title, 0.8, 0.45, 11.7, 0.8, 26)
        box(page, body, 0.85, 1.7, 11.5, 4.85, body_size)
        box(page, f"KUANGUARD 合成示範　{snapshot['snapshot_id']}　{len(deck.slides)}", 0.8, 6.92, 11.8, 0.3, 9, "516274")
        page.notes_slide.notes_text_frame.text = f"Snapshot {snapshot['snapshot_hash']}\n{DISCLOSURE}"
        return page
    slide(f"KUANGUARD {snapshot['service_code']} 檢測摘要", f"{snapshot['project_name']}\n版本 {snapshot['report_version']}\n{DISCLOSURE}", 24)
    stats = snapshot["statistics"]
    slide("檢測範圍與統計口徑", f"預定資產 {stats['planned_asset_count']}　完成 {stats['completed_asset_count']}\n弱點種類 {stats['definition_count']}　實例 {stats['occurrence_count']}\n受影響資產 {stats['affected_asset_count']}\nInformational {stats['informational_count']} 與 Unknown {stats['unknown_count']} 另列\n未檢測與證據不足不視為安全", 24)
    slide("四級弱點實例", "\n".join(f"{LEVEL_ZH[level]} {level}　{stats['severity_counts'][level]}" for level in RISK_LEVELS), 24)
    for level in ("Critical", "High"):
        group = [f for f in snapshot["findings"] if f["severity"] == level and (snapshot["service_code"] != "SHC" or f.get("check_status") == "failed")]
        if not group:
            slide(f"{LEVEL_ZH[level]}風險結果", "N/A")
        for finding in group:
            endpoint = finding["asset"]
            if finding.get("port") is not None:
                endpoint = (f"[{endpoint}]" if ":" in endpoint and "://" not in endpoint else endpoint) + f":{finding['port']}"
            body = f"來源 {finding['source_id']}\n資產 {endpoint}"
            if finding.get("location"):
                body += "\n位置 " + finding["location"]
            body += "\n說明 " + (finding["description"] or "來源未提供") + "\n改善 " + (finding["solution"] or "來源未提供")
            for index, chunk in enumerate(_text_chunks(body, 220), 1):
                slide(f"{LEVEL_ZH[level]}風險 {finding['source_id']}" + (f" 續頁 {index}" if index > 1 else ""), chunk, 24)
    if "shc" in stats:
        slide("健診符合狀態", "\n".join(f"{STATUS_ZH[key]}　{value}" for key, value in stats["shc"]["check_counts"].items()), 24)
    if snapshot["retest"]:
        body = "\n".join(f"{r['source_id']}　{RETEST_ZH[r['status']]}" for r in snapshot["retest"])
        for chunk in _text_chunks(body, 200):
            slide("複測追蹤", chunk)
    slide("資料追溯與限制", f"覆核 {snapshot['reviewer_id']}\n時間 {snapshot['reviewed_at']}\nParser {snapshot['parser_version']}\n模板 {snapshot['template_version']}\n完整來源與快照 SHA256 列於 manifest.json", 18)
    deck.core_properties.title = f"KUANGUARD {snapshot['service_code']} 檢測摘要"
    deck.core_properties.subject = snapshot["snapshot_id"]
    deck.save(path)


def generate_bundle(snapshot: dict, output_dir: Path) -> dict:
    """Atomically complete a local bundle; repeat calls verify immutable artifacts."""
    _verify_snapshot(snapshot)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    if manifest_path.exists():
        old = json.loads(manifest_path.read_text(encoding="utf-8"))
        if old.get("snapshot_hash") != snapshot["snapshot_hash"]:
            raise FileExistsError("IMMUTABLE_BUNDLE: 既有 bundle 屬於不同快照")
        for item in old["artifacts"]:
            artifact = output_dir / item["filename"]
            if not artifact.is_file() or hashlib.sha256(artifact.read_bytes()).hexdigest() != item["sha256"]:
                raise ValueError("ARTIFACT_CHECKSUM_MISMATCH: 既有 bundle 檔案被更動")
        return old
    if any((output_dir / name).exists() for name in ARTIFACT_NAMES):
        raise FileExistsError("INCOMPLETE_BUNDLE: 輸出目錄已有未完成產物，請建立新 job 目錄")
    temporary = Path(tempfile.mkdtemp(prefix="kg_report_", dir=output_dir.parent))
    try:
        _write_csv(snapshot, temporary / "all.csv")
        _write_xlsx(snapshot, temporary / "summary.xlsx")
        _write_docx(snapshot, temporary / "report.docx")
        fonts = _write_pdf(snapshot, temporary / "report.pdf")
        _write_pptx(snapshot, temporary / "summary.pptx")
        artifacts = [{"format": Path(name).suffix[1:], "filename": name, "bytes": (temporary / name).stat().st_size,
                      "sha256": hashlib.sha256((temporary / name).read_bytes()).hexdigest()} for name in ARTIFACT_NAMES]
        manifest = {"schema_version": "kuanguard.bundle/1", "snapshot_id": snapshot["snapshot_id"], "snapshot_hash": snapshot["snapshot_hash"],
                    "batch_id": snapshot["batch_id"], "report_version": snapshot["report_version"], "trace_id": snapshot["trace_id"],
                    "parser_version": snapshot["parser_version"], "generator_version": REPORT_VERSION, "template_version": snapshot["template_version"],
                    "rule_version": snapshot["rule_version"], "reviewer_id": snapshot["reviewer_id"], "reviewed_at": snapshot["reviewed_at"],
                    "sources": snapshot["sources"], "statistics": snapshot["statistics"], "synthetic": True, "disclosure": DISCLOSURE,
                    "artifacts": artifacts, "qa": {"structural": "generated_and_reopened", "fonts": fonts,
                                                  "visual": "requires_per_page_review", "office_render": "not_performed_by_generator",
                                                  "pdf_engine": "reportlab_native_pdf_not_office_conversion", "production_ready": False,
                                                  "production_blockers": snapshot["publication_gates"]["production_blockers"]}}
        # Reopen the actual files before committing the manifest, not just checking existence.
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pypdf import PdfReader
        Document(temporary / "report.docx")
        workbook = load_workbook(temporary / "summary.xlsx", read_only=True, data_only=False)
        if workbook["Findings"].max_row != len(snapshot["findings"]) + 1:
            raise ValueError("XLSX_ROW_COUNT_MISMATCH")
        workbook.close()
        if len(Presentation(temporary / "summary.pptx").slides) < 5:
            raise ValueError("PPTX_INCOMPLETE")
        pdf = PdfReader(temporary / "report.pdf")
        if not pdf.pages or snapshot["snapshot_id"] not in "".join(page.extract_text() or "" for page in pdf.pages):
            raise ValueError("PDF_SNAPSHOT_MISSING")
        manifest["qa"]["pdf_pages"] = len(pdf.pages)
        (temporary / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        for name in (*ARTIFACT_NAMES, "manifest.json"):
            os.replace(temporary / name, output_dir / name)
        return manifest
    finally:
        # Only remove the unique temporary directory created by this invocation.
        shutil.rmtree(temporary)
